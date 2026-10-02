#!/usr/bin/env python3
"""Build one PPTX slide deck per section from the [SLIDE n: title] cues in lecture scripts.

Usage
-----
    python voice-ai-agents-course/09-production/tools/slide_builder.py --course voice-ai-agents-course
    python voice-ai-agents-course/09-production/tools/slide_builder.py --course agent-observability-course --section 06
    python voice-ai-agents-course/09-production/tools/slide_builder.py --course . --scripts-dir 02-course-content
    python voice-ai-agents-course/09-production/tools/slide_builder.py --render-diagrams voice-ai-agents-course/10-graphics/diagrams

What it reads
-------------
``<course>/02-lecture-scripts/section-*.md`` (or ``--scripts-dir``). For every ``## Lecture x.y`` header it
collects the ``[SLIDE n: title]`` cues and the block beneath each cue: bullet lists, markdown tables,
``Table: a | b`` tables with ``- x | y`` rows, fenced code, ``Diagram:`` descriptions, block quotes and a
``Footer:`` line. The narration that follows a cue (up to the next visual cue) becomes the speaker notes.

What it writes
--------------
``<course>/10-graphics/slides/section-NN.pptx`` (or ``--out DIR``), styled per ``10-graphics/design-system.md``:
K2 title card per lecture, one slide per cue (K3 teaching / K4 diagram / K5 code / table), K6 recap cards
from ``[SLIDE n: Recap]`` cues, "You can now" cards, and a K7 next-up card from the lecture transition.

Diagrams: a cue that names a master diagram (``D3``) or whose title/description matches a diagram in
``<course>/10-graphics/diagrams/index.json`` gets that SVG rendered to PNG (cached in ``_preview/``). Any other
``Diagram:`` description is drawn as an amber dashed box flagged "DIAGRAM TO BUILD".

Rendering uses Playwright + Chromium (``PLAYWRIGHT_BROWSERS_PATH``), falling back to cairosvg. Without either,
diagrams fall back to the description box. Requires ``python-pptx`` (see requirements.txt).
"""
from __future__ import annotations

import argparse
import glob
import json
import os
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

# --------------------------------------------------------------------------------------------
# Design tokens (10-graphics/design-system.md)
# --------------------------------------------------------------------------------------------
NAVY = "0A1628"
NAVY_LIGHT = "1A2742"
NAVY_MID = "0F1E35"
TEAL = "00D4AA"
RED = "FF4B4B"
AMBER = "FFB020"
GRAY = "94A3B8"
GRAY_DARK = "64748B"
GRAY_LIGHT = "CBD5E1"
WHITE = "FFFFFF"
HEAD_FONT = "Inter"
BODY_FONT = "Inter"
MONO_FONT = "JetBrains Mono"

MAX_BULLETS = 5
TITLE_BAND = 180  # px of the 1920x1080 diagram canvas holding the diagram title (cropped on K3 slides)
MAX_CODE_LINES = 15
MAX_TABLE_ROWS = 8  # body rows per slide; longer tables continue on the next slide

VISUAL_CUES = ("SLIDE", "SCREEN", "CODE", "DEMO", "B-ROLL")

# --------------------------------------------------------------------------------------------
# Data model
# --------------------------------------------------------------------------------------------


@dataclass
class Slide:
    number: str
    title: str
    lecture_id: str = ""
    bullets: list[str] = field(default_factory=list)
    table: list[list[str]] = field(default_factory=list)  # first row = header
    code: str = ""
    code_lang: str = ""
    diagram: str = ""  # free-text diagram description
    text: list[str] = field(default_factory=list)  # prose / quotes shown on the slide
    footer: str = ""
    notes: str = ""
    line: int = 0

    @property
    def kind(self) -> str:
        t = self.title.strip().lower()
        if re.match(r"^recap\b", t):
            return "recap"
        if t.startswith("you can now"):
            return "youcannow"
        if self.code:
            return "code"
        if self.table:
            return "table"
        if self.diagram:
            return "diagram"
        return "teaching"


@dataclass
class Lecture:
    id: str
    title: str
    objective: str = ""
    recap: str = ""
    transition: str = ""
    slides: list[Slide] = field(default_factory=list)


@dataclass
class Section:
    number: int
    title: str
    path: str = ""
    lectures: list[Lecture] = field(default_factory=list)


# --------------------------------------------------------------------------------------------
# Parser
# --------------------------------------------------------------------------------------------

SECTION_RE = re.compile(r"^#\s+(?:Section|Module)\s+(\d+)\s*[:.\-—–]\s*(.+?)\s*$", re.I)
LECTURE_RE = re.compile(r"^##\s+Lecture\s+(\d+\.\d+[a-z]?)\b\s*(?:[—–:\-]\s*)?(.*?)\s*$")
SLIDE_RE = re.compile(r"^\[SLIDE\s+([0-9]+[a-z]?)\s*:\s*(.+)\]\s*$")
CUE_RE = re.compile(r"^\[(SLIDE|SCREEN|CODE|DEMO|B-ROLL|AVATAR|PAUSE)\b")
BULLET_RE = re.compile(r"^\s*(?:[-*•]|\d+[.)])\s+(.*)$")
DIAGRAM_START_RE = re.compile(
    r"^(diagram\b|timeline\b|stacked bar|two stacked bars|chart\b|waterfall\b|matrix\b|build\s*\d*\s*:)", re.I
)
TABLE_COLON_RE = re.compile(r"^Table\s*:\s*(.+\|.+)$", re.I)
FOOTER_RE = re.compile(r"^Footer\s*:\s*(.+)$", re.I)


def _strip_md(s: str) -> str:
    """Remove bold/italic markers but keep `code` backticks."""
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"(?<![\w*])\*(?!\s)(.+?)(?<!\s)\*(?![\w*])", r"\1", s)
    s = re.sub(r"(?<!\w)_(?!\s)(.+?)(?<!\s)_(?!\w)", r"\1", s)
    return s.strip()


def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [_strip_md(c.strip()) for c in line.split("|")]


def _is_separator_row(line: str) -> bool:
    return bool(re.match(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$", line.strip()))


def _clean_narration(text: str) -> str:
    text = re.sub(r"\[PAUSE[^\]]*\]", " ", text)
    text = re.sub(r"\[AVATAR\]", " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _first_sentence(text: str, max_words: int = 22) -> str:
    text = re.sub(r"^\s*1[.)]\s*", "", text.strip())
    m = re.split(r"(?<=[.!?])\s+(?=\d+[.)]\s|[A-Z])", text, maxsplit=1)
    s = m[0].strip()
    words = s.split()
    if len(words) > max_words:
        s = " ".join(words[:max_words]).rstrip(",;:") + "…"
    return s


def _blocks(lines: list[str], start: int) -> Iterable[tuple[int, int]]:
    """Yield (begin, end) index ranges of blank-line separated blocks starting at ``start``.
    Fenced code blocks are kept whole even if they contain blank lines."""
    i = start
    n = len(lines)
    while i < n:
        while i < n and not lines[i].strip():
            i += 1
        if i >= n:
            return
        b = i
        in_fence = False
        while i < n:
            s = lines[i].strip()
            if s.startswith("```"):
                in_fence = not in_fence
            elif not s and not in_fence:
                break
            i += 1
        yield b, i


def _block_is_body(block: list[str], first: bool, prev_ends_colon: bool) -> bool:
    head = block[0].strip()
    if CUE_RE.match(head) or head.startswith("#") or re.match(r"^\*\*(Recap|Transition)", head):
        return False
    if head.startswith("```") or head.startswith("|") or head.startswith(">"):
        return True
    if BULLET_RE.match(head) or TABLE_COLON_RE.match(head) or FOOTER_RE.match(head):
        return True
    if DIAGRAM_START_RE.match(head):
        return True
    return first or prev_ends_colon


def parse_slide_body(slide: Slide, body: list[str]) -> None:
    """Fill a Slide from the raw lines under its cue."""
    i = 0
    n = len(body)
    while i < n:
        raw = body[i]
        s = raw.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            lang = s[3:].strip()
            j = i + 1
            code: list[str] = []
            while j < n and not body[j].strip().startswith("```"):
                code.append(body[j].rstrip("\n"))
                j += 1
            slide.code = "\n".join(code)
            slide.code_lang = lang
            i = j + 1
            continue
        m = TABLE_COLON_RE.match(s)
        if m:
            slide.table = [_split_row(m.group(1))]
            j = i + 1
            while j < n and BULLET_RE.match(body[j]) and "|" in body[j]:
                slide.table.append(_split_row(BULLET_RE.match(body[j]).group(1)))
                j += 1
            i = j
            continue
        if s.startswith("|"):
            rows = []
            j = i
            while j < n and body[j].strip().startswith("|"):
                if not _is_separator_row(body[j]):
                    rows.append(_split_row(body[j]))
                j += 1
            slide.table = rows
            i = j
            continue
        m = FOOTER_RE.match(s)
        if m:
            slide.footer = _strip_md(m.group(1).strip().strip('"'))
            i += 1
            continue
        m = BULLET_RE.match(raw)
        if m:
            slide.bullets.append(_strip_md(m.group(1)))
            i += 1
            continue
        if s.startswith(">"):
            slide.text.append(_strip_md(s.lstrip(">").strip()))
            i += 1
            continue
        if DIAGRAM_START_RE.match(s) or slide.diagram:
            # Diagram description plus any prose lines that continue it.
            slide.diagram = (slide.diagram + " " + _strip_md(s)).strip()
            i += 1
            continue
        slide.text.append(_strip_md(s))
        i += 1


def parse_script(text: str, path: str = "") -> Section:
    lines = text.splitlines()
    section = Section(number=0, title="", path=path)
    m = re.search(r"section-(\d+)", os.path.basename(path))
    if m:
        section.number = int(m.group(1))
    lecture: Lecture | None = None
    in_fence = False
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            in_fence = not in_fence
            i += 1
            continue
        if in_fence:
            i += 1
            continue
        m = SECTION_RE.match(s)
        if m and not section.title:
            section.number = section.number or int(m.group(1))
            section.title = m.group(2).strip()
            i += 1
            continue
        m = LECTURE_RE.match(s)
        if m:
            lecture = Lecture(id=m.group(1), title=_strip_md(m.group(2)))
            section.lectures.append(lecture)
            i += 1
            continue
        if lecture is not None:
            # Field table rows and objective lists
            fm = re.match(r"^\|\s*(Learning objectives|One idea)\s*\|\s*(.+?)\s*\|\s*$", s, re.I)
            if fm:
                val = _strip_md(fm.group(2))
                if fm.group(1).lower() == "one idea" or not lecture.objective:
                    lecture.objective = _first_sentence(val)
                i += 1
                continue
            if re.match(r"^\*\*Learning objectives\*\*", s) and not lecture.objective:
                j = i + 1
                while j < n and not lines[j].strip():
                    j += 1
                if j < n and BULLET_RE.match(lines[j]):
                    lecture.objective = _first_sentence(BULLET_RE.match(lines[j]).group(1))
                i = j
                continue
            rm = re.match(r"^\*\*(Recap|Transition):\*\*\s*(.+)$", s)
            if rm:
                setattr(lecture, rm.group(1).lower(), _strip_md(rm.group(2)))
                i += 1
                continue
            hm = re.match(r"^###\s+(Recap|Transition)\s*$", s)
            if hm:
                j = i + 1
                para: list[str] = []
                while j < n and not lines[j].strip():
                    j += 1
                while j < n and lines[j].strip() and not lines[j].startswith("#"):
                    para.append(lines[j].strip())
                    j += 1
                setattr(lecture, hm.group(1).lower(), _strip_md(" ".join(para)))
                i = j
                continue
            sm = SLIDE_RE.match(s)
            if sm:
                slide = Slide(number=sm.group(1), title=_strip_md(sm.group(2)), lecture_id=lecture.id, line=i + 1)
                body: list[str] = []
                end = i + 1
                first = True
                prev_colon = False
                for b, e in _blocks(lines, i + 1):
                    block = lines[b:e]
                    if not _block_is_body(block, first, prev_colon):
                        break
                    body.extend(block + [""])
                    prev_colon = block[-1].rstrip().endswith(":")
                    first = False
                    end = e
                parse_slide_body(slide, body)
                # Narration: from end of body to next visual cue / header / recap marker.
                notes: list[str] = []
                j = end
                fence = False
                while j < n:
                    t = lines[j].strip()
                    if t.startswith("```"):
                        fence = not fence
                        j += 1
                        continue
                    if fence:
                        j += 1
                        continue
                    cm = CUE_RE.match(t)
                    if (cm and cm.group(1) == "SLIDE") or t.startswith("#") or t == "---":
                        break
                    if re.match(r"^\*\*(Recap|Transition):\*\*", t):
                        break
                    if cm and cm.group(1) in VISUAL_CUES:
                        # keep other visual cues so the editor sees where the slide cuts away
                        notes.append(t if len(t) <= 140 else t[:137].rstrip() + "...]")
                    else:
                        notes.append("" if (cm or not t) else t)  # [AVATAR]/[PAUSE] and blanks = breaks
                    j += 1
                slide.notes = _clean_narration(re.sub(r"\n{2,}", "\n\n", "\n".join(notes)).strip())
                lecture.slides.append(slide)
                i = end
                continue
        i += 1
    return section


# --------------------------------------------------------------------------------------------
# Diagram index and matching
# --------------------------------------------------------------------------------------------

EXPLICIT_D_RE = re.compile(r"\*{0,2}\bD(\d{1,2})\b\*{0,2}(?:\s*\(?\s*(?:build|step)\s*(\d))?", re.I)


@dataclass
class DiagramRef:
    id: str
    svg: Path
    step: int | None = None


def load_index(diagram_dir: Path) -> dict:
    p = diagram_dir / "index.json"
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        return {}


def _diagram_file(diagram_dir: Path, entry: dict, step: int | None) -> Path:
    if step and str(step) in (entry.get("steps") or {}):
        return diagram_dir / entry["steps"][str(step)]["file"]
    return diagram_dir / entry["file"]


def match_diagram(slide: Slide, index: dict, diagram_dir: Path) -> DiagramRef | None:
    """Return the master diagram a slide should show, or None."""
    if not index:
        return None
    diagrams = index.get("diagrams", {})
    title = slide.title
    desc = slide.diagram
    # 1) explicit ID in the cue title or the diagram description ("D3", "**D3** build 2")
    for src in (title, desc):
        for m in EXPLICIT_D_RE.finditer(src):
            did = f"D{int(m.group(1))}"
            if did in diagrams:
                step = int(m.group(2)) if m.group(2) else None
                return DiagramRef(did, _diagram_file(diagram_dir, diagrams[did], step), step)
    # 2) keyword match, only for slides that look like a diagram or whose title names one
    hay = f"{title} {desc}".lower()
    best: tuple[int, str] | None = None
    for did, entry in diagrams.items():
        kws = [k.lower() for k in entry.get("keywords", [])]
        hits = [k for k in kws if k in hay]
        if not hits:
            continue
        in_lecture = slide.lecture_id in entry.get("lectures", [])
        title_hit = any(k in title.lower() for k in kws)
        if not (desc or title_hit):
            continue
        if not in_lecture and not desc:
            continue
        score = 3 * len(hits) + (4 if in_lecture else 0) + (2 if title_hit else 0)
        if best is None or score > best[0]:
            best = (score, did)
    if best is None or best[0] < 5:
        return None
    did = best[1]
    entry = diagrams[did]
    step = None
    for st, spec in sorted((entry.get("steps") or {}).items(), key=lambda kv: int(kv[0])):
        if any(k.lower() in hay for k in spec.get("keywords", [])):
            step = int(st)
    return DiagramRef(did, _diagram_file(diagram_dir, entry, step), step)


# --------------------------------------------------------------------------------------------
# SVG → PNG rendering (Playwright Chromium, fallback cairosvg)
# --------------------------------------------------------------------------------------------


def _chromium_executable() -> str | None:
    base = os.environ.get("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
    for pat in ("chromium-*/chrome-linux*/chrome", "chromium-*/chrome-mac*/Chromium.app/Contents/MacOS/Chromium",
                "chromium-*/chrome-win*/chrome.exe", "chromium_headless_shell-*/chrome-*/headless_shell"):
        hits = sorted(glob.glob(os.path.join(base, pat)))
        if hits:
            return hits[-1]
    return None


class Renderer:
    """Renders SVG files to 1920x1080 PNGs. Use as a context manager so one browser serves many files."""

    def __init__(self) -> None:
        self._pw = None
        self._browser = None
        self._page = None
        self.backend = None

    def __enter__(self) -> "Renderer":
        try:
            from playwright.sync_api import sync_playwright  # type: ignore

            self._pw = sync_playwright().start()
            try:
                self._browser = self._pw.chromium.launch()
            except Exception:
                exe = _chromium_executable()
                if not exe:
                    raise
                self._browser = self._pw.chromium.launch(executable_path=exe)
            self._page = self._browser.new_page(viewport={"width": 1920, "height": 1080})
            self.backend = "playwright"
        except Exception:
            self._close_pw()
            try:
                import cairosvg  # type: ignore  # noqa: F401

                self.backend = "cairosvg"
            except Exception:
                self.backend = None
        return self

    def _close_pw(self) -> None:
        try:
            if self._browser:
                self._browser.close()
            if self._pw:
                self._pw.stop()
        except Exception:
            pass
        self._browser = self._pw = self._page = None

    def __exit__(self, *exc) -> None:
        self._close_pw()

    def render(self, svg: Path, png: Path, hide_title: bool = False) -> bool:
        png.parent.mkdir(parents=True, exist_ok=True)
        svg_text = svg.read_text(encoding="utf-8")
        if hide_title:
            svg_text = svg_text.replace("</svg>", "<style>#diagram-title{display:none}</style></svg>", 1)
        if self.backend == "playwright":
            html = (
                "<!doctype html><html><head><meta charset='utf-8'><style>html,body{margin:0;padding:0;"
                "background:#0A1628}svg{display:block;width:1920px;height:1080px}</style></head><body>"
                + svg_text
                + "</body></html>"
            )
            self._page.set_content(html, wait_until="load")
            self._page.evaluate("document.fonts.ready")
            clip = {"x": 0, "y": TITLE_BAND, "width": 1920, "height": 1080 - TITLE_BAND} if hide_title else \
                {"x": 0, "y": 0, "width": 1920, "height": 1080}
            self._page.screenshot(path=str(png), clip=clip)
        elif self.backend == "cairosvg":
            import cairosvg  # type: ignore

            if hide_title:  # crop the title band by shifting the viewBox
                svg_text = svg_text.replace('viewBox="0 0 1920 1080" width="1920" height="1080"',
                                            f'viewBox="0 {TITLE_BAND} 1920 {1080 - TITLE_BAND}" width="1920" '
                                            f'height="{1080 - TITLE_BAND}"', 1)
            cairosvg.svg2png(bytestring=svg_text.encode("utf-8"), write_to=str(png))
        else:
            return False
        _compress_png(png)
        return True


def _image_aspect(png: Path) -> float:
    try:
        from PIL import Image  # type: ignore

        with Image.open(png) as im:
            return im.width / im.height
    except Exception:
        return 16 / 9


def _compress_png(png: Path) -> None:
    """Quantize to a 256-colour palette: diagrams are flat colour, this cuts size ~3x."""
    try:
        from PIL import Image  # type: ignore

        im = Image.open(png).convert("RGB")
        im.quantize(colors=256, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(
            png, optimize=True)
    except Exception:
        pass


def cached_png(renderer: Renderer | None, svg: Path, hide_title: bool) -> Path | None:
    if not svg.exists():
        return None
    suffix = "-notitle" if hide_title else ""
    png = svg.parent / "_preview" / f"{svg.stem}{suffix}.png"
    if png.exists() and png.stat().st_mtime >= svg.stat().st_mtime:
        return png
    if renderer is None or renderer.backend is None:
        return None
    return png if renderer.render(svg, png, hide_title=hide_title) else None


def render_directory(diagram_dir: Path) -> list[Path]:
    out = []
    with Renderer() as r:
        if r.backend is None:
            print("No renderer available (install playwright or cairosvg).", file=sys.stderr)
            return out
        for svg in sorted(diagram_dir.glob("*.svg")):
            png = diagram_dir / "_preview" / f"{svg.stem}.png"
            if r.render(svg, png):
                out.append(png)
    return out


# --------------------------------------------------------------------------------------------
# PPTX writer
# --------------------------------------------------------------------------------------------


def _pptx():
    try:
        import pptx  # type: ignore  # noqa: F401
    except ImportError:  # pragma: no cover
        sys.exit("python-pptx is required: pip install -r voice-ai-agents-course/09-production/tools/requirements.txt")
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    return Presentation, RGBColor, MSO_LINE, MSO_SHAPE, MSO_ANCHOR, PP_ALIGN, Emu, Inches, Pt


class DeckWriter:
    """Writes the slides for one section. Sizes follow the design system's px values on a 1920x1080 canvas
    (13.333 in wide, so 1 pt = 2 px): headings 48 px = 24 pt, body 24 px = 12 pt is too small for slides
    watched on a phone, so body text uses the Callout size (28 px = 14 pt) and up."""

    def __init__(self, course_label: str, section: Section) -> None:
        (self.Presentation, self.RGB, self.MSO_LINE, self.MSO_SHAPE, self.MSO_ANCHOR, self.PP_ALIGN,
         self.Emu, self.Inches, self.Pt) = _pptx()
        self.prs = self.Presentation()
        self.prs.slide_width = self.Inches(13.333)
        self.prs.slide_height = self.Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.course_label = course_label
        self.section = section
        self.stats = {"slides": 0, "diagrams": 0, "to_build": 0, "recaps": 0, "title_cards": 0}

    # -- primitives -------------------------------------------------------------------------
    def rgb(self, hexstr: str):
        return self.RGB.from_string(hexstr)

    def new_slide(self, lecture_id: str = "", notes: str = ""):
        s = self.prs.slides.add_slide(self.blank)
        fill = s.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb(NAVY)
        if lecture_id:
            self.footer(s, f"{self.course_label} · Section {self.section.number:02d} · Lecture {lecture_id}")
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        self.stats["slides"] += 1
        return s

    def textbox(self, slide, x, y, w, h, wrap=True, anchor="top"):
        tb = slide.shapes.add_textbox(self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.margin_left = tf.margin_right = self.Inches(0.05)
        tf.margin_top = tf.margin_bottom = self.Inches(0.03)
        tf.vertical_anchor = {"top": self.MSO_ANCHOR.TOP, "middle": self.MSO_ANCHOR.MIDDLE,
                              "bottom": self.MSO_ANCHOR.BOTTOM}[anchor]
        return tb, tf

    def runs(self, paragraph, text: str, size: float, color: str = WHITE, bold: bool = False,
             font: str = BODY_FONT, mono_color: str = TEAL) -> None:
        """Add text to a paragraph, rendering `backtick` spans in JetBrains Mono."""
        parts = re.split(r"(`[^`]+`)", text)
        for part in parts:
            if not part:
                continue
            r = paragraph.add_run()
            if part.startswith("`") and part.endswith("`") and len(part) > 1:
                r.text = part[1:-1]
                r.font.name = MONO_FONT
                r.font.size = self.Pt(size * 0.92)
                r.font.color.rgb = self.rgb(mono_color)
                r.font.bold = False
            else:
                r.text = part
                r.font.name = font
                r.font.size = self.Pt(size)
                r.font.color.rgb = self.rgb(color)
                r.font.bold = bold

    def bullet_xml(self, paragraph, color: str = TEAL, char: str = "•", indent_in: float = 0.32) -> None:
        from pptx.oxml.ns import qn
        from lxml import etree

        pPr = paragraph._p.get_or_add_pPr()
        pPr.set("marL", str(int(self.Inches(indent_in))))
        pPr.set("indent", str(-int(self.Inches(indent_in))))
        for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        srgb = etree.SubElement(buClr, qn("a:srgbClr"))
        srgb.set("val", color)
        buFont = etree.SubElement(pPr, qn("a:buFont"))
        buFont.set("typeface", HEAD_FONT)
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", char)

    def rect(self, slide, x, y, w, h, fill: str | None = None, line: str | None = None, line_w: float = 2,
             dash: bool = False, rounded: bool = True):
        shp = slide.shapes.add_shape(self.MSO_SHAPE.ROUNDED_RECTANGLE if rounded else self.MSO_SHAPE.RECTANGLE,
                                     self.Inches(x), self.Inches(y), self.Inches(w), self.Inches(h))
        if rounded:
            # 8 px radius on a 1920 canvas = 0.056 in; adjustment is a fraction of the short side
            shp.adjustments[0] = min(0.5, 0.056 / max(0.01, min(w, h)))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = self.rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = self.rgb(line)
            shp.line.width = self.Pt(line_w * 0.75)
            if dash:
                shp.line.dash_style = self.MSO_LINE.DASH
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        shp.text_frame.text = ""
        return shp

    def footer(self, slide, text: str) -> None:
        _, tf = self.textbox(slide, 0.5, 7.02, 9.5, 0.35)
        self.runs(tf.paragraphs[0], text, 9, GRAY_DARK)

    def title(self, slide, text: str, y: float = 0.45, size: float = 26, width: float = 12.3) -> float:
        _, tf = self.textbox(slide, 0.5, y, width, 0.95, anchor="top")
        self.runs(tf.paragraphs[0], text, size, WHITE, bold=True, font=HEAD_FONT)
        tf.paragraphs[0].line_spacing = 1.1
        lines = 1 + (len(text) * size * 0.0086 > width)  # conservative: fallback fonts run wider than Inter
        bar_y = y + (0.55 if lines == 1 else 1.0) * size / 26
        self.rect(slide, 0.55, bar_y + 0.12, 0.9, 0.06, fill=TEAL, rounded=False)
        return bar_y + 0.45

    def bullets(self, slide, items: list[str], x, y, w, h, size: float = 18) -> None:
        _, tf = self.textbox(slide, x, y, w, h)
        for k, item in enumerate(items):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            self.bullet_xml(p)
            p.line_spacing = 1.2
            p.space_after = self.Pt(size * 0.7)
            self.runs(p, item, size, WHITE)

    def prose(self, slide, lines: list[str], x, y, w, h, size: float = 16, color: str = GRAY_LIGHT) -> None:
        _, tf = self.textbox(slide, x, y, w, h)
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = self.Pt(6)
            self.runs(p, line, size, color)

    def footnote(self, slide, text: str) -> None:
        _, tf = self.textbox(slide, 0.5, 6.62, 12.3, 0.4)
        self.runs(tf.paragraphs[0], text, 11, GRAY)

    # -- scene kits -------------------------------------------------------------------------
    def k2_title_card(self, lec: Lecture, notes: str = "") -> None:
        s = self.new_slide(lec.id, notes)
        _, tf = self.textbox(s, 7.3, 0.35, 5.55, 0.6)
        tf.paragraphs[0].alignment = self.PP_ALIGN.RIGHT
        self.runs(tf.paragraphs[0], f"SECTION {self.section.number} · {self.section.title.upper()}", 10, GRAY)
        _, tf = self.textbox(s, 1.0, 1.55, 11.3, 0.5)
        self.runs(tf.paragraphs[0], f"LECTURE {lec.id}", 12, TEAL, bold=True)
        _, tf = self.textbox(s, 1.0, 2.05, 11.3, 1.8, anchor="top")
        self.runs(tf.paragraphs[0], lec.title, 30, WHITE, bold=True, font=HEAD_FONT)
        tf.paragraphs[0].line_spacing = 1.1
        if lec.objective:
            self.rect(s, 1.0, 4.15, 0.07, 1.25, fill=TEAL, rounded=False)
            _, tf = self.textbox(s, 1.3, 4.1, 10.6, 1.4, anchor="middle")
            self.runs(tf.paragraphs[0], lec.objective, 16, GRAY_LIGHT)
            tf.paragraphs[0].line_spacing = 1.3
        self.stats["title_cards"] += 1

    def k6_card(self, heading: str, items: list[str], lec_id: str, notes: str, label: str = "RECAP") -> None:
        s = self.new_slide(lec_id, notes)
        _, tf = self.textbox(s, 1.5, 0.9, 10.3, 0.5)
        tf.paragraphs[0].alignment = self.PP_ALIGN.CENTER
        self.runs(tf.paragraphs[0], label, 12, TEAL, bold=True)
        _, tf = self.textbox(s, 1.5, 1.35, 10.3, 0.9)
        tf.paragraphs[0].alignment = self.PP_ALIGN.CENTER
        self.runs(tf.paragraphs[0], heading, 26, WHITE, bold=True, font=HEAD_FONT)
        items = items[:5] if items else []
        top = 2.75 if len(items) <= 3 else 2.45
        for k, item in enumerate(items):
            y = top + k * 1.05
            self.rect(s, 2.2, y, 8.9, 0.85, fill=NAVY_LIGHT, line=None)
            _, tf = self.textbox(s, 2.4, y, 0.6, 0.85, anchor="middle")
            self.runs(tf.paragraphs[0], "✓", 22, TEAL, bold=True)
            _, tf = self.textbox(s, 3.05, y, 7.9, 0.85, anchor="middle")
            self.runs(tf.paragraphs[0], item, 18, WHITE, bold=True)
        self.stats["recaps"] += 1

    def k7_next_up(self, next_title: str, teaser: str, lec_id: str) -> None:
        s = self.new_slide(lec_id)
        _, tf = self.textbox(s, 2.6, 2.5, 1.2, 1.2, anchor="middle")
        self.runs(tf.paragraphs[0], "→", 54, TEAL, bold=True)
        _, tf = self.textbox(s, 3.9, 2.35, 8.0, 0.5)
        self.runs(tf.paragraphs[0], "Next up:", 14, GRAY)
        _, tf = self.textbox(s, 3.9, 2.8, 8.0, 1.2)
        self.runs(tf.paragraphs[0], next_title, 26, WHITE, bold=True, font=HEAD_FONT)
        if teaser:
            _, tf = self.textbox(s, 3.9, 4.05, 8.0, 1.2)
            self.runs(tf.paragraphs[0], teaser, 14, GRAY)

    def table(self, slide, rows: list[list[str]], x, y, w, h_max) -> None:
        ncols = max(len(r) for r in rows)
        rows = [r + [""] * (ncols - len(r)) for r in rows]
        longest = max(len(c) for r in rows for c in r) if rows else 10
        size = 14 if longest < 45 and len(rows) <= 6 else 12 if len(rows) <= 8 else 11
        row_h = 0.48 if size >= 14 else 0.42
        gt = slide.shapes.add_table(len(rows), ncols, self.Inches(x), self.Inches(y), self.Inches(w),
                                    self.Inches(min(h_max, row_h * len(rows))))
        tbl = gt.table
        # column widths proportional to content length (bounded)
        lens = [max(6, min(60, max(len(rows[r][c]) for r in range(len(rows))))) for c in range(ncols)]
        total = sum(lens)
        for c in range(ncols):
            tbl.columns[c].width = self.Inches(w * lens[c] / total)
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = tbl.cell(r, c)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb(TEAL if r == 0 else (NAVY_LIGHT if r % 2 else NAVY_MID))
                cell.margin_left = cell.margin_right = self.Inches(0.1)
                cell.margin_top = cell.margin_bottom = self.Inches(0.05)
                cell.vertical_anchor = self.MSO_ANCHOR.MIDDLE
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.text = ""
                if r == 0:
                    self.runs(p, val, size, NAVY, bold=True, mono_color=NAVY)
                else:
                    color = WHITE
                    self.runs(p, val, size, color, bold=(c == 0))

    def code_block(self, slide, code: str, label: str, x, y, w, h) -> None:
        lines = code.rstrip().splitlines()
        hidden = 0
        if len(lines) > MAX_CODE_LINES:
            hidden = len(lines) - (MAX_CODE_LINES - 1)
            lines = lines[: MAX_CODE_LINES - 1] + [f"# ... {hidden} more lines in the script"]
        self.rect(slide, x, y, w, h, fill=NAVY_MID, line=NAVY_LIGHT)
        self.rect(slide, x, y, 0.06, h, fill=TEAL, rounded=False)
        if label:
            _, tf = self.textbox(slide, x + 0.25, y + 0.08, w - 0.4, 0.35)
            self.runs(tf.paragraphs[0], label, 10, GRAY, font=MONO_FONT)
        longest = max((len(l) for l in lines), default=20)
        size = 14 if longest <= 70 and len(lines) <= 12 else 12 if longest <= 90 else 10
        _, tf = self.textbox(slide, x + 0.25, y + (0.45 if label else 0.2), w - 0.4, h - 0.5, wrap=False)
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line if line else " "
            r.font.name = MONO_FONT
            r.font.size = self.Pt(size)
            is_comment = line.lstrip().startswith(("#", "//"))
            r.font.color.rgb = self.rgb(GRAY if is_comment else WHITE)

    def diagram_to_build(self, slide, desc: str, x, y, w, h) -> None:
        self.rect(slide, x, y, w, h, fill=NAVY_MID, line=AMBER, dash=True)
        _, tf = self.textbox(slide, x + 0.3, y + 0.2, w - 0.6, 0.4)
        self.runs(tf.paragraphs[0], "DIAGRAM TO BUILD", 11, AMBER, bold=True)
        _, tf = self.textbox(slide, x + 0.3, y + 0.65, w - 0.6, h - 0.8)
        words = desc.split()
        if len(words) > 120:
            desc = " ".join(words[:120]) + " …"
        self.runs(tf.paragraphs[0], desc, 13 if len(words) > 60 else 15, GRAY_LIGHT)
        tf.paragraphs[0].line_spacing = 1.25
        self.stats["to_build"] += 1

    # -- slide from a cue -------------------------------------------------------------------
    def content_slide(self, sl: Slide, diagram_png: Path | None, diagram_full_png: Path | None) -> None:
        kind = sl.kind
        if kind in ("recap", "youcannow"):
            items = sl.bullets or [t for t in sl.text if t]
            if kind == "youcannow" and len(items) == 1 and "·" in items[0]:
                items = [p.strip() for p in items[0].split("·") if p.strip()]
            heading = sl.title if kind == "youcannow" else "Key takeaways"
            label = "SECTION COMPLETE" if kind == "youcannow" else "RECAP"
            self.k6_card(heading, items, sl.lecture_id, sl.notes, label=label)
            return

        has_text = bool(sl.bullets or sl.text)
        if diagram_full_png is not None and len(sl.bullets) > MAX_BULLETS and not sl.table and not sl.code:
            # the bullets describe the diagram's boxes: show the diagram, keep the text for the editor
            sl = Slide(**{**sl.__dict__, "bullets": [], "text": [],
                          "notes": sl.notes + "\n\nSlide text (drawn on the diagram): " + " · ".join(sl.bullets)})
            has_text = False
        # Full-bleed K4: a matched diagram with nothing else to say
        if diagram_full_png is not None and not has_text and not sl.table and not sl.code:
            s = self.new_slide("", sl.notes)
            s.shapes.add_picture(str(diagram_full_png), 0, 0, width=self.prs.slide_width)
            # lecture ID top right: the diagram's own footnote sits bottom left
            _, tf = self.textbox(s, 6.8, 0.12, 6.35, 0.3)
            tf.paragraphs[0].alignment = self.PP_ALIGN.RIGHT
            self.runs(tf.paragraphs[0], f"Section {self.section.number:02d} · Lecture {sl.lecture_id} · {sl.title}",
                      9, GRAY_DARK)
            if sl.footer:
                self.footnote(s, sl.footer)
            self.stats["diagrams"] += 1
            return

        bullets = sl.bullets
        chunks = [bullets[i:i + MAX_BULLETS] for i in range(0, len(bullets), MAX_BULLETS)] or [[]]
        table_chunks: list[list[list[str]]] = []
        if sl.table:
            head, body = sl.table[0], sl.table[1:]
            table_chunks = [[head] + body[i:i + MAX_TABLE_ROWS] for i in range(0, max(1, len(body)), MAX_TABLE_ROWS)]
        n_pages = max(len(chunks), len(table_chunks), 1)
        for page in range(n_pages):
            title = sl.title + (f" ({page + 1}/{n_pages})" if n_pages > 1 else "")
            s = self.new_slide(sl.lecture_id, sl.notes if page == 0 else "(continued) " + sl.notes[:200])
            top = self.title(s, title, size=24 if len(title) > 60 else 26)
            items = chunks[page] if page < len(chunks) else []
            right_visual = diagram_png is not None or (sl.diagram and not sl.table and not sl.code)
            text_lines = sl.text if page == 0 else []
            bottom = 6.55 if sl.footer else 6.9
            if sl.code and page == 0:
                if items or text_lines:
                    self.bullets(s, items, 0.5, top, 4.4, bottom - top, size=16) if items else \
                        self.prose(s, text_lines, 0.5, top, 4.4, bottom - top)
                    self.code_block(s, sl.code, sl.code_lang, 5.1, top, 7.75, bottom - top)
                else:
                    self.code_block(s, sl.code, sl.code_lang, 0.5, top, 12.35, bottom - top)
            elif table_chunks and page < len(table_chunks):
                y = top
                if text_lines or items:
                    self.prose(s, text_lines + items, 0.5, y, 12.3, 0.8, size=14)
                    y += 0.75
                self.table(s, table_chunks[page], 0.5, y, 12.35, bottom - y)
            elif right_visual and page == 0:
                if items or text_lines:
                    if items:
                        self.bullets(s, items, 0.5, top, 4.75, bottom - top, size=16)
                    else:
                        self.prose(s, text_lines, 0.5, top, 4.75, bottom - top, size=15)
                    vx, vw = 5.45, 7.4
                else:
                    vx, vw = 0.9, 11.55
                aspect = _image_aspect(diagram_png) if diagram_png is not None else 16 / 9
                box_w = vw
                vh = vw / aspect
                if vh > bottom - top:
                    vh = bottom - top
                    vw = vh * aspect
                    vx = vx + (box_w - vw) / 2
                vy = top + max(0.0, (bottom - top - vh) / 2)
                if diagram_png is not None:
                    pic = s.shapes.add_picture(str(diagram_png), self.Inches(vx), self.Inches(vy),
                                               width=self.Inches(vw))
                    pic.line.color.rgb = self.rgb(NAVY_LIGHT)
                    self.stats["diagrams"] += 1
                else:
                    self.diagram_to_build(s, sl.diagram, vx, top, vw, bottom - top)
            else:
                if items:
                    self.bullets(s, items, 0.7, top + 0.1, 11.9, bottom - top - 0.1,
                                 size=22 if sum(len(i) for i in items) < 220 else 18)
                if text_lines:
                    y = top + (0.1 if not items else min(bottom - 1.2, top + 0.75 * len(items) + 0.6))
                    big = not items and sum(len(t) for t in text_lines) < 160
                    if big:
                        self.rect(s, 0.7, y + 0.2, 0.07, 1.2, fill=TEAL, rounded=False)
                    self.prose(s, text_lines, 1.0 if big else 0.7, y + (0.25 if big else 0), 11.6,
                               bottom - y, size=22 if big else 15, color=WHITE if big else GRAY_LIGHT)
            if sl.footer:
                self.footnote(s, sl.footer)

    def save(self, path: Path) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))


# --------------------------------------------------------------------------------------------
# Orchestration
# --------------------------------------------------------------------------------------------

COURSE_LABELS = {
    "voice-ai-agents-course": "Production Voice AI Agents",
    "agent-observability-course": "AI Agent Observability & Cost Control",
}


def course_label(course_dir: Path) -> str:
    name = course_dir.resolve().name
    return COURSE_LABELS.get(name, "AI Agent Testing & Evaluation")


def build_section(section: Section, course_dir: Path, out_dir: Path, renderer: Renderer | None,
                  index: dict, diagram_dir: Path, next_title: str = "") -> tuple[Path, dict]:
    w = DeckWriter(course_label(course_dir), section)
    for li, lec in enumerate(section.lectures):
        w.k2_title_card(lec, notes=f"Title card (K2). {lec.title}.")
        for sl in lec.slides:
            ref = match_diagram(sl, index, diagram_dir) if sl.kind not in ("recap", "youcannow") else None
            png = full = None
            if ref is not None:
                has_text = bool(sl.bullets or sl.text) and len(sl.bullets) <= MAX_BULLETS
                if has_text or sl.table or sl.code:
                    png = cached_png(renderer, ref.svg, hide_title=True)
                else:
                    full = cached_png(renderer, ref.svg, hide_title=False)
            w.content_slide(sl, png, full)
        nxt = section.lectures[li + 1].title if li + 1 < len(section.lectures) else next_title
        if nxt and lec.transition:
            w.k7_next_up(nxt, _first_sentence(lec.transition, 18), lec.id)
    out = out_dir / f"section-{section.number:02d}.pptx"
    w.save(out)
    return out, w.stats


def find_scripts(course_dir: Path, scripts_dir: str | None) -> list[Path]:
    base = (course_dir / scripts_dir) if scripts_dir else course_dir / "02-lecture-scripts"
    if scripts_dir and Path(scripts_dir).is_absolute():
        base = Path(scripts_dir)
    return sorted(base.glob("section-*.md"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--course", help="Course directory (e.g. voice-ai-agents-course, or . for Course 2)")
    ap.add_argument("--section", help="Only this section number, e.g. 03")
    ap.add_argument("--out", help="Output directory (default <course>/10-graphics/slides)")
    ap.add_argument("--scripts-dir", help="Scripts folder relative to --course (default 02-lecture-scripts)")
    ap.add_argument("--diagrams-dir", help="Diagram folder (default <course>/10-graphics/diagrams)")
    ap.add_argument("--no-render", action="store_true", help="Do not render SVGs; use cached PNGs only")
    ap.add_argument("--render-diagrams", metavar="DIR", help="Render every SVG in DIR to DIR/_preview and exit")
    args = ap.parse_args(argv)

    if args.render_diagrams:
        pngs = render_directory(Path(args.render_diagrams))
        print(f"Rendered {len(pngs)} PNGs into {Path(args.render_diagrams) / '_preview'}")
        return 0 if pngs else 1
    if not args.course:
        ap.error("--course is required")

    course_dir = Path(args.course)
    scripts = find_scripts(course_dir, args.scripts_dir)
    if args.section:
        scripts = [p for p in scripts if re.search(rf"section-0*{int(args.section)}\b", p.name)]
    if not scripts:
        print(f"No section-*.md scripts found for {course_dir}", file=sys.stderr)
        return 1
    out_dir = Path(args.out) if args.out else course_dir / "10-graphics" / "slides"
    diagram_dir = Path(args.diagrams_dir) if args.diagrams_dir else course_dir / "10-graphics" / "diagrams"
    index = load_index(diagram_dir)

    sections = [parse_script(p.read_text(encoding="utf-8"), str(p)) for p in scripts]
    all_sections = sections
    if args.section:
        all_sections = [parse_script(p.read_text(encoding="utf-8"), str(p)) for p in find_scripts(course_dir, args.scripts_dir)]

    def next_lecture_title(sec: Section) -> str:
        later = [s for s in all_sections if s.number > sec.number and s.lectures]
        return later[0].lectures[0].title if later else ""

    renderer_cm = Renderer() if not args.no_render else None
    renderer = renderer_cm.__enter__() if renderer_cm else None
    try:
        total = 0
        for sec in sections:
            out, stats = build_section(sec, course_dir, out_dir, renderer, index, diagram_dir, next_lecture_title(sec))
            size_kb = out.stat().st_size / 1024
            total += out.stat().st_size
            print(f"{out}  {stats['slides']:3d} slides  {stats['diagrams']:2d} diagrams  "
                  f"{stats['to_build']:2d} to build  {stats['recaps']:2d} recap/you-can-now  {size_kb:,.0f} KB")
        print(f"Total {total / 1024 / 1024:.1f} MB in {out_dir}")
    finally:
        if renderer_cm:
            renderer_cm.__exit__(None, None, None)
    return 0


if __name__ == "__main__":
    sys.exit(main())
